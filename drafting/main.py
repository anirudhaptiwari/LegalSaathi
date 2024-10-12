import streamlit as st
from docx import Document
from docx.shared import Pt
from docx.enum.style import WD_STYLE_TYPE
import os
import json
import mammoth
import re
from docx2pdf import convert
import io
import tempfile

# Function to replace text in paragraphs and runs
def replace_text_in_paragraph(paragraph, old_text, new_text):
    if old_text in paragraph.text:
        inline = paragraph.runs
        for i in range(len(inline)):
            if old_text in inline[i].text:
                text = inline[i].text.replace(old_text, new_text)
                inline[i].text = text

# Function to generate the document
def generate_document(selected_contract, form_details, local_file_path):
    doc = Document(local_file_path)

    # Create styles if they don't exist
    styles = doc.styles
    if 'Heading 1' not in styles:
        styles.add_style('Heading 1', WD_STYLE_TYPE.PARAGRAPH)
        styles['Heading 1'].font.size = Pt(16)
        styles['Heading 1'].font.bold = True
    if 'Heading 2' not in styles:
        styles.add_style('Heading 2', WD_STYLE_TYPE.PARAGRAPH)
        styles['Heading 2'].font.size = Pt(14)
        styles['Heading 2'].font.bold = True

    # Apply formatting and replacements
    for paragraph in doc.paragraphs:
        # Apply heading styles
        if paragraph.text.startswith("SERVICE AGREEMENT") or paragraph.text.startswith("NON-DISCLOSURE AGREEMENT"):
            paragraph.style = styles['Heading 1']
        elif paragraph.text.strip().startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.", "10.", "11.", "12.")):
            paragraph.style = styles['Heading 2']
        
        # Replace placeholders
        for placeholder, value in form_details.items():
            pattern = rf'\[.*?:\s*{re.escape(placeholder)}\]'
            matches = re.findall(pattern, paragraph.text)
            for match in matches:
                replace_text_in_paragraph(paragraph, match, value)

    return doc

# CSS for styling the preview
preview_css = """
<style>
    .contract-preview {
        background-color: white;
        color: black;
        font-family: Arial, sans-serif;
        padding: 20px;
        border: 1px solid #ddd;
        border-radius: 5px;
    }
    .contract-preview h1 {
        color: #2c3e50;
        font-size: 24px;
        margin-bottom: 20px;
    }
    .contract-preview h2 {
        color: #34495e;
        font-size: 20px;
        margin-top: 15px;
        margin-bottom: 10px;
    }
    .contract-preview p {
        margin-bottom: 10px;
        line-height: 1.5;
    }
</style>
"""

def main():
    # Ensure the docs directory exists
    docs_dir = os.path.join(os.path.dirname(__file__), '..', 'docs')
    os.makedirs(docs_dir, exist_ok=True)

    # Load contract types from JSON file
    with open(os.path.join(os.path.dirname(__file__), 'contract_types.json'), 'r') as f:
        contract_types = json.load(f)

    # Load placeholder questions from JSON file
    with open(os.path.join(os.path.dirname(__file__), 'placeholder_questions.json'), 'r') as f:
        placeholder_questions = json.load(f)

    # Sidebar: Contract Types
    st.sidebar.title("Corporate & Business Contracts")
    selected_contract = st.sidebar.selectbox("Select a Contract Type", options=list(contract_types.keys()))

    # Main Content: Display Contract Form
    if selected_contract:
        st.header(f"{selected_contract} Generator")
        st.subheader("Please fill in the following details:")

        # Get the document template path
        template_path = contract_types[selected_contract]
        local_file_path = os.path.join(docs_dir, template_path)

        # Create input fields for each placeholder
        form_details = {}
        for placeholder, question in placeholder_questions[selected_contract].items():
            form_details[placeholder] = st.text_input(question)

        # Button to generate preview and enable downloads
        if st.button("Generate Contract"):
            st.write("Generating contract...")

            # Generate the document
            doc = generate_document(selected_contract, form_details, local_file_path)

            # Save the document to a temporary file
            with tempfile.NamedTemporaryFile(delete=False, suffix='.docx') as tmp_file:
                doc.save(tmp_file.name)
                tmp_file_path = tmp_file.name

            # Convert the document to HTML for preview
            with open(tmp_file_path, 'rb') as docx_file:
                result = mammoth.convert_to_html(docx_file)
                html_content = result.value

            # Wrap the HTML content with our custom CSS
            styled_html = f"{preview_css}<div class='contract-preview'>{html_content}</div>"

            # Display the preview
            st.subheader("Contract Preview")
            st.components.v1.html(styled_html, height=600, scrolling=True)

            # Save as DOCX
            output_docx_path = os.path.join(docs_dir, f'{selected_contract.lower().replace(" ", "_")}_output.docx')
            doc.save(output_docx_path)

            # Convert to PDF
            output_pdf_path = output_docx_path.replace('.docx', '.pdf')
            convert(output_docx_path, output_pdf_path)

            # Download buttons
            st.subheader("Download Options")
            
            # DOCX download
            with open(output_docx_path, 'rb') as f:
                docx_bytes = f.read()
            
            st.download_button(
                label="Download Contract (DOCX)",
                data=docx_bytes,
                file_name=f'{selected_contract.lower().replace(" ", "_")}.docx',
                mime='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            )

            # PDF download
            with open(output_pdf_path, 'rb') as f:
                pdf_bytes = f.read()
            
            st.download_button(
                label="Download Contract (PDF)",
                data=pdf_bytes,
                file_name=f'{selected_contract.lower().replace(" ", "_")}.pdf',
                mime='application/pdf'
            )

            # Remove the temporary file
            os.unlink(tmp_file_path)

        # Display legal compliance notice
        st.info("This contract generator is designed to comply with Indian laws. However, it is recommended to have the final contract reviewed by a legal professional to ensure full compliance with current regulations and your specific business needs.")

if __name__ == "__main__":
    main()
